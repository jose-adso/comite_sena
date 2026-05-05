import os
import json
from datetime import datetime, timedelta
from flask import render_template, request, redirect, url_for, flash, session, jsonify
from flask import send_file
from models.database import db
from models.grabacion import Grabacion
from models.reclamo import Reclamo
from models.falla import Falla
from routes.auth import auth_bp
from routes.utils import es_admin, etiqueta_rol_visible, get_rol_activo, es_superadmin
import tempfile

# Whisper para transcripciÃ³n local (offline)
try:
    import whisper
    WHISPER_AVAILABLE = True
    _whisper_model = None
except ImportError:
    WHISPER_AVAILABLE = False
    whisper = None

import subprocess
import shutil

def convertir_audio_a_wav(input_path):
    """Convierte audio webm a wav usando ffmpeg directamente"""
    try:
        ffmpeg_path = r"C:\Users\JOSEROJAS\AppData\Local\Microsoft\WinGet\Packages\Gyan.FFmpeg_Microsoft.Winget.Source_8wekyb3d8bbwe\ffmpeg-8.1-full_build\bin\ffmpeg.exe"
        
        # Verificar que ffmpeg existe
        if not os.path.exists(ffmpeg_path):
            ffmpeg_path = "ffmpeg"  # Usar el del sistema si no existe el especÃ­fico
        
        wav_path = input_path.replace('.webm', '.wav')
        
        # Convertir usando ffmpeg directamente
        result = subprocess.run(
            [ffmpeg_path, "-i", input_path, "-acodec", "pcm_s16le", "-ar", "16000", "-ac", "1", wav_path, "-y"],
            capture_output=True,
            text=True,
            timeout=120
        )
        
        if result.returncode != 0:
            print(f"FFmpeg error: {result.stderr}")
            return input_path
        
        # Eliminar archivo original
        if os.path.exists(input_path):
            os.unlink(input_path)
        
        return wav_path
    except Exception as e:
        print(f"Error convirtiendo audio: {e}")
        return input_path

def get_whisper_model():
    """Carga el modelo Whisper lazy (una sola vez)"""
    global _whisper_model
    if not WHISPER_AVAILABLE:
        return None
    if _whisper_model is None:
        try:
            print("Cargando modelo Whisper 'small'...")
            _whisper_model = whisper.load_model("small")
            print("Modelo Whisper cargado correctamente.")
        except Exception as e:
            print(f"Error cargando modelo Whisper: {e}")
            return None
    return _whisper_model


@auth_bp.route('/grabar-reunion')
def vista_grabar_reunion():
    """Vista para grabar reuniones"""
    if 'usuario_id' not in session and 'username' not in session:
        flash('Debe iniciar sesiÃ³n primero', 'error')
        return redirect(url_for('auth.index'))
    
    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        flash('No tiene permisos para acceder a esta sección', 'error')
        return redirect(url_for('auth.dashboard'))
    
    today = datetime.now().date().isoformat()
    
    grabaciones = Grabacion.query.filter_by(
        usuario_id=session.get('usuario_id')
    ).order_by(Grabacion.fecha_registro.desc()).all()
    
    grabaciones_formatted = []
    for g in grabaciones:
        duracion_str = f"{g.duracion // 60:02d}:{g.duracion % 60:02d}"
        grabaciones_formatted.append({
            'id': g.id,
            'titulo': g.titulo,
            'fecha': g.fecha.strftime('%Y-%m-%d') if g.fecha else '',
            'duracion': duracion_str
        })
    
    reclamos = Reclamo.query.order_by(Reclamo.fecha_registro.desc()).limit(50).all()
    reclamos_formatted = [{
        'id': r.id,
        'titulo': f'#{r.id} - {r.nombre_aprendiz} ({r.tipo_reclamo})'
    } for r in reclamos]
    
    fallas = Falla.query.order_by(Falla.fecha_registro.desc()).limit(50).all()
    fallas_formatted = [{
        'id': -f.id,
        'titulo': f'Falla #{f.id} - {f.nombre_aprendiz} ({f.numero_ficha})'
    } for f in fallas]
    
    return render_template(
        'grabar_reunion.html',
        username=session.get('username'),
        rol_visible=etiqueta_rol_visible(),
        today_date=today,
        grabaciones=grabaciones_formatted,
        reclamos=reclamos_formatted + fallas_formatted
    )


@auth_bp.route('/guardar-grabacion', methods=['POST'])
def guardar_grabacion():
    """Guardar una grabaciÃ³n de audio"""
    if 'usuario_id' not in session and 'username' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'})

    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        return jsonify({'success': False, 'error': 'No tiene permisos'})

    try:
        titulo = request.form.get('titulo', '').strip()
        claim_id = request.form.get('claim_id', type=int)
        fecha_str = request.form.get('fecha', '').strip()
        duracion = request.form.get('duracion', 0, type=int)
        transcripcion = request.form.get('transcripcion', '').strip()
        hora_inicio_str = request.form.get('hora_inicio', '')
        audio = request.files.get('audio')

        if not titulo:
            return jsonify({'success': False, 'error': 'El tÃ­tulo es requerido'})

        if not audio:
            return jsonify({'success': False, 'error': 'El audio es requerido'})

        fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date() if fecha_str else datetime.now().date()

        uploads_dir = os.path.join(os.path.dirname(__file__), '..', 'static', 'grabaciones')
        os.makedirs(uploads_dir, exist_ok=True)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'grabacion_{timestamp}.webm'
        filepath = os.path.join(uploads_dir, filename)

        audio.save(filepath)

        real_claim_id = claim_id if claim_id and claim_id > 0 else None
        real_falla_id = -claim_id if claim_id and claim_id < 0 else None

        print(f'[DEBUG GUARDAR] claim_id recibido: {claim_id}, real_claim_id: {real_claim_id}, real_falla_id: {real_falla_id}')

        hora_inicio = None
        if hora_inicio_str:
            try:
                hora_inicio = datetime.fromisoformat(hora_inicio_str.replace('Z', '+00:00'))
            except:
                hora_inicio = datetime.now()

        grabacion = Grabacion(
            titulo=titulo,
            fecha=fecha,
            duracion=duracion,
            hora_inicio=hora_inicio,
            audio_path=f'grabaciones/{filename}',
            transcripcion=transcripcion or None,
            claim_id=real_claim_id,
            falla_id=real_falla_id,
            usuario_id=session.get('usuario_id')
        )
        db.session.add(grabacion)
        db.session.commit()

        return jsonify({'success': True})

    except Exception as e:
        print(f'Error al guardar grabaciÃ³n: {e}')
        return jsonify({'success': False, 'error': str(e)})


@auth_bp.route('/editar-grabacion/<int:grabacion_id>')
def editar_grabacion(grabacion_id):
    """Editar/transcribir una grabaciÃ³n"""
    if 'usuario_id' not in session and 'username' not in session:
        flash('Debe iniciar sesiÃ³n primero', 'error')
        return redirect(url_for('auth.index'))
    
    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        flash('No tiene permisos', 'error')
        return redirect(url_for('auth.dashboard'))
    
    grabacion = Grabacion.query.get_or_404(grabacion_id)
    
    if request.method == 'POST':
        transcripcion = request.form.get('transcripcion', '').strip()
        grabacion.transcripcion = transcripcion
        db.session.commit()
        flash('TranscripciÃ³n guardada', 'success')
        return redirect(url_for('auth.vista_grabar_reunion'))
    
    return render_template(
        'editar_grabacion.html',
        username=session.get('username'),
        rol_visible=etiqueta_rol_visible(),
        grabacion=grabacion
    )


@auth_bp.route('/eliminar-grabacion/<int:grabacion_id>', methods=['POST'])
def eliminar_grabacion(grabacion_id):
    """Eliminar una grabaciÃ³n"""
    if 'usuario_id' not in session and 'username' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'})
    
    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        return jsonify({'success': False, 'error': 'No tiene permisos'})
    
    try:
        grabacion = Grabacion.query.get_or_404(grabacion_id)
        
        # Eliminar archivo de audio si existe
        if grabacion.audio_path:
            audio_file = os.path.join(os.path.dirname(__file__), '..', 'static', grabacion.audio_path)
            if os.path.exists(audio_file):
                os.unlink(audio_file)
        
        # Eliminar de la base de datos
        db.session.delete(grabacion)
        db.session.commit()
        
        return jsonify({'success': True})
    
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})




# AlmacÃ©n temporal de transcripciones en vivo por usuario (en memoria)
_transcripciones_en_vivo = {}

@auth_bp.route('/transcribir-audio', methods=['POST'])
def transcribir_audio():
    """Transcribe un archivo de audio usando Whisper (offline)"""
    if 'usuario_id' not in session and 'username' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'})
    
    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        return jsonify({'success': False, 'error': 'No tiene permisos'})
    
    # Verificar si Whisper está disponible
    if not WHISPER_AVAILABLE:
        return jsonify({'success': False, 'error': 'Whisper no instalado. Ejecuta: pip install openai-whisper'})
    
    audio_file = request.files.get('audio')
    if not audio_file:
        return jsonify({'success': False, 'error': 'No se recibiÃ³ archivo de audio'})
    
    try:
        # Guardar temporalmente el audio
        with tempfile.NamedTemporaryFile(delete=False, suffix='.webm') as tmp:
            audio_file.save(tmp.name)
            tmp_path = tmp.name
        
        # Convertir webm a wav
        tmp_path = convertir_audio_a_wav(tmp_path)
        
        # Cargar modelo y transcribir
        model = get_whisper_model()
        if model is None:
            return jsonify({'success': False, 'error': 'No se pudo cargar el modelo Whisper'})
        
        # Transcribir en espaÃ±ol
        result = model.transcribe(tmp_path, language="es", fp16=False)
        texto = result["text"]
        
        # Limpiar archivo temporal
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        
        return jsonify({'success': True, 'text': texto.strip()})
    
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@auth_bp.route('/transcribir-en-vivo', methods=['POST'])
def transcribir_en_vivo():
    """Recibe fragmentos de audio y devuelve transcripciÃ³n acumulada usando Whisper"""
    if 'usuario_id' not in session and 'username' not in session:
        return jsonify({'success': False, 'error': 'No autorizado'})
    
    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        return jsonify({'success': False, 'error': 'No tiene permisos'})
    
    if not WHISPER_AVAILABLE:
        return jsonify({'success': False, 'error': 'Whisper no disponible'})
    
    audio_file = request.files.get('audio')
    if not audio_file:
        return jsonify({'success': False, 'error': 'No se recibiÃ³ audio'})
    
    uid = session.get('usuario_id', 0)
    
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix='.webm') as tmp:
            audio_file.save(tmp.name)
            tmp_path = tmp.name
        
        tmp_path = convertir_audio_a_wav(tmp_path)
        
        model = get_whisper_model()
        if model is None:
            return jsonify({'success': False, 'text': _transcripciones_en_vivo.get(uid, '')})
        
        result = model.transcribe(tmp_path, language="es", fp16=False)
        nuevo_texto = result["text"].strip()
        
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        
        # Acumular texto
        prev = _transcripciones_en_vivo.get(uid, '')
        if nuevo_texto and nuevo_texto not in prev:
            # Evitar duplicados al final
            if prev and prev.endswith(nuevo_texto[:30]):
                _transcripciones_en_vivo[uid] = prev + nuevo_texto[30:]
            else:
                _transcripciones_en_vivo[uid] = (prev + ' ' + nuevo_texto).strip()
        
        return jsonify({'success': True, 'text': _transcripciones_en_vivo.get(uid, '')})
    
    except Exception as e:
        print('Error transcribir-en-vivo:', e)
        return jsonify({'success': False, 'text': _transcripciones_en_vivo.get(uid, '')})


@auth_bp.route('/limpiar-transcripcion', methods=['POST'])
def limpiar_transcripcion():
    """Limpia la transcripciÃ³n en vivo almacenada para el usuario"""
    uid = session.get('usuario_id', 0)
    _transcripciones_en_vivo[uid] = ''
    return jsonify({'success': True})


def reemplazar_en_documento(doc, reemplazos):
    """Reemplaza marcadores en párrafos, tablas y runs de un documento docx"""
    def _reemplazar_en_texto(texto):
        resultado = texto
        for marca, valor in reemplazos.items():
            if marca in resultado:
                resultado = resultado.replace(marca, str(valor))
        return resultado

    # Párrafos del cuerpo
    for para in doc.paragraphs:
        texto_original = para.text
        texto_nuevo = _reemplazar_en_texto(texto_original)
        if texto_nuevo != texto_original:
            # Limpiar runs y poner texto completo en el primer run
            para.clear()
            para.add_run(texto_nuevo)

    # Tablas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    texto_original = para.text
                    texto_nuevo = _reemplazar_en_texto(texto_original)
                    if texto_nuevo != texto_original:
                        para.clear()
                        para.add_run(texto_nuevo)


@auth_bp.route('/generar-acta/<int:grabacion_id>')
def generar_acta(grabacion_id):
    """Genera un acta en formato .docx basada en la plantilla ACTA"""
    if 'usuario_id' not in session and 'username' not in session:
        flash('Debe iniciar sesiÃ³n primero', 'error')
        return redirect(url_for('auth.index'))

    rol = get_rol_activo()
    if rol not in ['planta', 'super admin'] and not es_superadmin():
        flash('No tiene permisos', 'error')
        return redirect(url_for('auth.dashboard'))
    
    grabacion = Grabacion.query.get_or_404(grabacion_id)

    try:
        from docx import Document
        import locale
        try:
            locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')
        except:
            try:
                locale.setlocale(locale.LC_TIME, 'Spanish_Spain.1252')
            except:
                pass

        plantilla_path = os.path.join(os.path.dirname(__file__), '..', 'static', 'plantillas', 'ACTA.docx')
        doc = Document(plantilla_path)

        # Fecha
        fecha_reunion = grabacion.fecha or datetime.now().date()
        dia = fecha_reunion.day
        mes = fecha_reunion.strftime('%B')
        ano = fecha_reunion.year
        fecha_acta_str = f"{dia} de {mes} de {ano}"

        # Hora de inicio (usar la guardada en BD o la actual)
        if grabacion.hora_inicio:
            hora_inicio_dt = grabacion.hora_inicio
        else:
            hora_inicio_dt = datetime.now()
        hora_inicio = hora_inicio_dt.strftime('%I:%M %p').lstrip('0').replace('AM', 'a.m.').replace('PM', 'p.m.')
        # Hora fin calculada con duraciÃ³n
        hora_fin_dt = hora_inicio_dt + timedelta(seconds=grabacion.duracion)
        hora_fin = hora_fin_dt.strftime('%I:%M %p').lstrip('0').replace('AM', 'a.m.').replace('PM', 'p.m.')

        # Datos del aprendiz e instructor (pueden venir de Reclamo o Falla)
        nombre_aprendiz = 'No especificado'
        nombre_ficha = 'No especificado'
        numero_ficha = 'No especificado'
        nombre_instructor = ''
        cedula_instructor = ''
        documento_aprendiz = ''
        correo_aprendiz = ''
        telefono_aprendiz = ''
        descripcion_faltas = ''

        # Obtener datos desde claim (reclamo) o falla
        origen = None
        print(f'[DEBUG ACTA] grabacion.claim_id={grabacion.claim_id}, grabacion.falla_id={grabacion.falla_id}')
        if grabacion.claim_id:
            print(f'[DEBUG ACTA] Buscando Reclamo id={grabacion.claim_id}')
            origen = Reclamo.query.get(grabacion.claim_id)
            if origen:
                print(f'[DEBUG ACTA] Reclamo encontrado: nombre_aprendiz="{origen.nombre_aprendiz}", nombre_ficha="{origen.nombre_ficha}", numero_ficha="{origen.numero_ficha}"')
            else:
                print(f'[DEBUG ACTA] Reclamo NO encontrado')
        elif grabacion.falla_id:
            print(f'[DEBUG ACTA] Buscando Falla id={grabacion.falla_id}')
            origen = Falla.query.get(grabacion.falla_id)
            if origen:
                print(f'[DEBUG ACTA] Falla encontrada: nombre_aprendiz="{origen.nombre_aprendiz}", nombre_ficha="{origen.nombre_ficha}", numero_ficha="{origen.numero_ficha}"')
            else:
                print(f'[DEBUG ACTA] Falla NO encontrada')

        if origen:
            nombre_aprendiz = origen.nombre_aprendiz or 'No especificado'
            nombre_ficha = origen.nombre_ficha or 'No especificado'
            numero_ficha = origen.numero_ficha or 'No especificado'
            nombre_instructor = origen.nombre_instructor or ''
            cedula_instructor = origen.cedula_instructor or ''
            documento_aprendiz = origen.documento_aprendiz or ''
            correo_aprendiz = origen.correo_aprendiz or ''
            telefono_aprendiz = origen.telefono_aprendiz or ''
            descripcion_faltas = origen.descripcion_faltas or ''

        transcripcion = grabacion.transcripcion or 'Sin transcripción'

        # Diccionario de reemplazos
        # Diccionario de reemplazos - placeholders EXACTOS de ACTA.docx
        reemplazos = {
            '[fecha acta]': fecha_acta_str,
            '[inicia]': hora_inicio,
            '[fin]': hora_fin,
            '[nombre aperdiz]': nombre_aprendiz,
            '[nombre ficha]': nombre_ficha,
            '[numero ficha]': numero_ficha,
            '[aquí se aplica la transcripcion]': transcripcion,
        }

        print(f'[DEBUG ACTA] Reemplazos: {reemplazos}')

        reemplazar_en_documento(doc, reemplazos)

        with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp:
            doc.save(tmp.name)
            tmp_path = tmp.name

        filename = f"acta_grabacion_{grabacion_id}.docx"
        return send_file(tmp_path, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    except Exception as e:
        print(f'Error generando acta: {e}')
        flash(f'Error al generar el acta: {str(e)}', 'error')
        return redirect(url_for('auth.vista_grabar_reunion'))

