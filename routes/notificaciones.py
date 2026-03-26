import os
import json
import tempfile
import uuid
from datetime import datetime, timedelta

from flask import render_template, request, redirect, url_for, flash, session
from docx import Document

from models.falla import Falla
from routes.auth import auth_bp
from routes.utils import es_admin, _replace_placeholders_doc, _enviar_pdf_siempre, etiqueta_rol_visible, get_rol_activo


# ── Helpers de notificaciones ─────────────────────────────────────────────────
def obtener_plantillas_notificacion():
    plantillas_dir = os.path.join(os.path.dirname(__file__), '..', 'static', 'plantillas')
    if not os.path.exists(plantillas_dir):
        return []
    return sorted([f for f in os.listdir(plantillas_dir) if f.lower().endswith('.docx')])


def formatear_fecha_larga(fecha_obj):
    meses = [
        'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
        'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
    ]
    return f"a los {fecha_obj.day} dias del mes de {meses[fecha_obj.month - 1]} del {fecha_obj.year}"


def formatear_fecha_citacion(fecha_obj):
    """Formatea fecha para citación: 27 de marzo del 2026"""
    meses = [
        'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
        'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
    ]
    return f"{fecha_obj.day} de {meses[fecha_obj.month - 1]} del {fecha_obj.year}"


def _ruta_horarios_citacion():
    return os.path.join(os.path.dirname(__file__), '..', 'instance', 'horarios_citacion.json')


def _bloques_horarios_disponibles():
    bloques = []
    actual = datetime.strptime('08:00', '%H:%M')
    fin = datetime.strptime('11:45', '%H:%M')
    while actual <= fin:
        bloques.append(actual.strftime('%H:%M'))
        actual += timedelta(minutes=15)
    actual = datetime.strptime('14:00', '%H:%M')
    fin = datetime.strptime('16:45', '%H:%M')
    while actual <= fin:
        bloques.append(actual.strftime('%H:%M'))
        actual += timedelta(minutes=15)
    return bloques


def _cargar_horarios_citacion():
    ruta = _ruta_horarios_citacion()
    if not os.path.exists(ruta):
        return {}
    try:
        with open(ruta, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _guardar_horarios_citacion(data):
    ruta = _ruta_horarios_citacion()
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def obtener_o_asignar_hora_citacion(fecha_citacion, falla_id):
    fecha_key = fecha_citacion.isoformat()
    falla_key = str(falla_id)
    data = _cargar_horarios_citacion()
    agenda_dia = data.get(fecha_key, {})
    if not isinstance(agenda_dia, dict):
        agenda_dia = {}
    if falla_key in agenda_dia:
        return agenda_dia[falla_key]
    ocupadas = set(str(h) for h in agenda_dia.values())
    for hora in _bloques_horarios_disponibles():
        if hora not in ocupadas:
            agenda_dia[falla_key] = hora
            data[fecha_key] = agenda_dia
            _guardar_horarios_citacion(data)
            return hora
    return None


# ── Rutas ─────────────────────────────────────────────────────────────────────
@auth_bp.route('/vista-notificaciones')
def vista_notificaciones():
    if 'usuario_id' not in session:
        flash('Debe iniciar sesión primero', 'error')
        return redirect(url_for('auth.index'))

    if not es_admin():
        flash('No tiene permisos para ver esta sección', 'error')
        return redirect(url_for('auth.dashboard'))

    filtro_instructor = request.args.get('filter_instructor', '').strip()
    filtro_aprendiz = request.args.get('filter_aprendiz', '').strip()

    query = Falla.query
    if filtro_instructor:
        query = query.filter(Falla.nombre_instructor.ilike(f'%{filtro_instructor}%'))
    if filtro_aprendiz:
        query = query.filter(Falla.nombre_aprendiz.ilike(f'%{filtro_aprendiz}%'))

    fallas = query.order_by(Falla.fecha_registro.desc()).all()
    plantillas = obtener_plantillas_notificacion()

    return render_template(
        'notificaciones.html',
        username=session['username'],
        rol=get_rol_activo(),
        rol_visible=etiqueta_rol_visible(),
        fallas=fallas,
        plantillas=plantillas,
        filtro_instructor=filtro_instructor,
        filtro_aprendiz=filtro_aprendiz,
    )


@auth_bp.route('/generar-notificacion-pdf', methods=['POST'])
def generar_notificacion_pdf():
    if 'usuario_id' not in session:
        flash('Debe iniciar sesión primero', 'error')
        return redirect(url_for('auth.index'))

    if not es_admin():
        flash('No tiene permisos para realizar esta acción', 'error')
        return redirect(url_for('auth.dashboard'))

    falla_id = request.form.get('falla_id', type=int)
    plantilla_nombre = (request.form.get('plantilla') or '').strip()
    fecha_citacion_raw = (request.form.get('fecha_citacion') or '').strip()

    if not falla_id or not plantilla_nombre:
        flash('Debe seleccionar una plantilla y un registro.', 'error')
        return redirect(url_for('auth.vista_notificaciones'))

    falla = Falla.query.get_or_404(falla_id)

    try:
        fecha_citacion = (
            datetime.strptime(fecha_citacion_raw, '%Y-%m-%d').date()
            if fecha_citacion_raw
            else datetime.now().date()
        )
        fecha_carta = datetime.now().date()
        hora_citacion = obtener_o_asignar_hora_citacion(fecha_citacion, falla.id)

        if not hora_citacion:
            flash('No hay horarios disponibles para la fecha seleccionada. Use otra fecha.', 'error')
            return redirect(url_for('auth.vista_notificaciones'))

        plantillas_dir = os.path.join(os.path.dirname(__file__), '..', 'static', 'plantillas')
        plantilla_path = os.path.join(plantillas_dir, plantilla_nombre)

        if not os.path.exists(plantilla_path):
            flash('La plantilla seleccionada no existe.', 'error')
            return redirect(url_for('auth.vista_notificaciones'))

        doc = Document(plantilla_path)
        replacements = {
            '[Fecha Carta]': formatear_fecha_larga(fecha_carta),
            '[Nombre Ficha]': falla.nombre_ficha or '',
            '[Numero Ficha]': falla.numero_ficha or '',
            '[NOMBRE FICHA]': falla.nombre_ficha or '',
            '[NUMERO FICHA]': falla.numero_ficha or '',
            '[Nombre instructor]': falla.nombre_instructor or '',
            '[Nombre Aprendiz]': falla.nombre_aprendiz or '',
            '[documento aprendiz]': falla.documento_aprendiz or '',
            '[Fecha Citación]': formatear_fecha_citacion(fecha_citacion),
            '[Hora]': hora_citacion,
            '[HORA]': hora_citacion,
        }
        _replace_placeholders_doc(doc, replacements)

        temp_dir = tempfile.gettempdir()
        temp_docx = os.path.join(temp_dir, f'notificacion_{falla.id}_{uuid.uuid4().hex}.docx')
        temp_pdf = os.path.join(temp_dir, f'notificacion_{falla.id}_{uuid.uuid4().hex}.pdf')
        doc.save(temp_docx)

        return _enviar_pdf_siempre(
            temp_docx=temp_docx,
            temp_pdf=temp_pdf,
            download_name=f'notificacion_{falla.numero_ficha}_{falla.nombre_aprendiz}.pdf',
            doc_para_respaldo=doc,
            titulo_respaldo='Notificacion generada',
        )
    except Exception as e:
        flash(f'Error al generar la notificación: {str(e)}', 'error')
        return redirect(url_for('auth.vista_notificaciones'))
