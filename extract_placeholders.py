import re
from docx import Document

doc = Document(r'C:\Users\JOSEROJAS\Documents\macro sena\llamado\templates\FORMATO COMITE.docx')

text_parts = []

# Paragraphs
for p in doc.paragraphs:
    text_parts.append(p.text)

# Tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            text_parts.append(cell.text)

# Headers and Footers
for section in doc.sections:
    for p in section.header.paragraphs:
        text_parts.append(p.text)
    for p in section.footer.paragraphs:
        text_parts.append(p.text)

full_text = "\n".join(text_parts)

# Find all placeholders: square brackets [...] or curly braces {...}
pattern = r'\[[^\]]*\]|\{[^{}]*\}'
matches = re.findall(pattern, full_text)

# Deduplicate while preserving order
seen = set()
unique_matches = []
for m in matches:
    if m not in seen:
        seen.add(m)
        unique_matches.append(m)

for m in unique_matches:
    print(m)
